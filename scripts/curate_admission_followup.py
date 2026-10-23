"""Read public primary sources into a quarantined evidence ledger; never fit models."""
from pathlib import Path
import argparse, hashlib, html, json, re

ROOT = Path(__file__).resolve().parents[1]
D = ROOT / 'docs/admission_followup_20261022'
AA = set('ACDEFGHIKLMNPQRSTVWY')

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def patent_tables(path, required):
    result = {}
    for table in re.findall(r'<table\b.*?</table>', path.read_text(encoding='utf-8'), re.S | re.I):
        text = re.sub(r'\s+', ' ', html.unescape(re.sub('<[^>]+>', ' ', table))).strip()
        match = re.match(r'\(SEQ ID NO:\s*(\d+)\)', text)
        if not match:
            continue
        sid = int(match[1])
        if sid not in required:
            continue
        sequence = re.sub(r'[\s.]', '', text[match.end():])
        if not sequence or not set(sequence) <= AA:
            raise ValueError('Invalid patent protein sequence')
        if sid in result and result[sid] != sequence:
            raise ValueError(f'Conflicting SEQ ID {sid}')
        result[sid] = sequence
    if set(result) != set(required):
        raise ValueError('Missing required patent sequence ID')
    return result

def st25(text, sid):
    match = re.search(r'<210>\s*' + str(sid) + r'\s*\n(.*?)(?=<210>|\Z)', text, re.S)
    if not match:
        raise ValueError(f'Missing ST25 SEQ {sid}')
    block = match[1]
    size = int(re.search(r'<211>\s*(\d+)', block)[1])
    body = re.split(r'<400>\s*' + str(sid) + r'\s*\n', block)[1]
    triples = 'Ala Arg Asn Asp Cys Gln Glu Gly His Ile Leu Lys Met Phe Pro Ser Thr Trp Tyr Val'.split()
    letters = 'ARNDCQEGHILKMFPSTWYV'
    mapping = dict(zip(triples, letters))
    sequence = ''.join(mapping[x] for x in re.findall(r'\b(?:' + '|'.join(triples) + r')\b', body))
    if len(sequence) != size:
        raise ValueError(f'ST25 length mismatch {sid}: {len(sequence)} != {size}')
    return sequence

def sequence_ledger():
    from pypdf import PdfReader
    old = patent_tables(D/'sources/US20220088190A1.html', [22,28,30,36,42,44,61,66,68])
    new = patent_tables(D/'sources/US20260083847A1.html', [49,54,56])
    pdf = PdfReader(D/'sources/KR20220082046A.pdf')
    text = '\n'.join(f'===PDF_PAGE_{i}===\n' + page.extract_text() for i,page in enumerate(pdf.pages,1))
    # Only explicitly published sequence IDs; no germline reconstruction or silent terminal repair.
    rows = []
    specs = [('A3V', old, 'US20220088190A1', 61, 66, 68, 'CALSEAGTYKYIF', 'CASSVAGGGQETQYF'),
             ('A11Va', old, 'US20220088190A1', 36, 42, 44, 'CAVNPPDTGFQKLVF', 'CASSLSFRQGLREQYF'),
             ('A11Vb', old, 'US20220088190A1', 22, 28, 30, 'CAVDKDGGYQKVTF', 'CSASPRAGQLSSYNSPLHF'),
             ('A11Vc', new, 'US20260083847A1', 49, 54, 56, 'CATDPGGFKTIF', 'CASSLYGGSISYEQYF')]
    for name, table, source, a, b, pair, ca, cb in specs:
        alpha, beta = table[a], table[b]
        if ca not in alpha or cb not in beta:
            raise ValueError('Paired CDR identity mismatch')
        # A paired cassette can alter the alpha terminal residue at the linker junction.
        if alpha.split('NIQNPDP')[0] not in table[pair] or beta.split('EDL')[0] not in table[pair]:
            raise ValueError('Published paired cassette does not support both chain sequences')
        rows.append(dict(receptor_name=name, pmid='39287991', alpha_published=alpha,
                         beta_published=beta, cdr3_alpha=ca, cdr3_beta=cb,
                         source=source, alpha_seq_id=a, beta_seq_id=b, paired_seq_id=pair,
                         sequence_scope='published full chains including leader and constant; no mature-domain trimming',
                         identity_bridge='paired CDR3 matches primary Table 1; terminal CDR conventions differ; exact experimental construct equivalence not established',
                         training_allowed=False, independent_acceptance_allowed=False))
    alpha = st25(text, 9)
    for name, b in [('A3-05', 256), ('A3-10', 261)]:
        beta = st25(text, b)
        rows.append(dict(receptor_name=name, pmid='36174557', alpha_published=alpha,
                         beta_published=beta, cdr3_alpha='CAVRPGGAGSYQLTF',
                         cdr3_beta='CASGINLASGQFF' if b == 256 else 'CASSKNLVDEQYF',
                         source='KR20220082046A; alias/pair mapping WO2021074249A1',
                         alpha_seq_id=9, beta_seq_id=b, source_pdf_pages=[109,110,179 if b == 256 else 183],
                         sequence_scope='published variable-domain entries include leaders; no mature-domain trimming',
                         identity_bridge='WO alias-to-sequence mapping plus KR same-family ST25 entries; not a clinical dosing construct',
                         training_allowed=False, independent_acceptance_allowed=False))
    for row in rows:
        if row['cdr3_alpha'] not in row['alpha_published'] or row['cdr3_beta'] not in row['beta_published']:
            raise ValueError('CDR absent from published chain')
    extracts = {str(sid): re.search(r'<210>\s*'+str(sid)+r'\s*\n(.*?)(?=<210>|\Z)',text,re.S)[0] for sid in [9,256,261]}
    return dict(receptors=rows, selected_ST25_text=extracts, training_release_created=False,
                caveats=['Research/public patent sequences are not clinical manufacturing DNA.',
                         'SEQ171 in US20220088190A1 is incomplete; SEQ26 omits a terminal residue relative to the full beta chain; neither was silently repaired.'])

def kras_cells():
    import openpyxl
    path=ROOT/'docs/measured_intake_20261021/source_metadata/kras_measurements.xlsx'
    book=openpyxl.load_workbook(path,data_only=True)
    library_path=ROOT/'docs/measured_intake_20261021/source_metadata/kras_peptides.xlsx'
    library_book=openpyxl.load_workbook(library_path,data_only=True)
    library={i:[] for i in range(10)};wt='VVVGAVGVGK'
    for source_row,record in enumerate(library_book['Supplemental File 2'].values,1):
        if source_row<=3:continue
        library_id,decorated,position,kras_position,declared_aa=record
        sequence=decorated.removeprefix('H-').removesuffix('-OH')
        p=position-1
        aa=sequence[p]
        if not (kras_position==position+6 and aa in AA and aa!=wt[p] and sequence==wt[:p]+aa+wt[p+1:]):
            raise ValueError('Primary library entry identity conflict')
        library[p].append(dict(id=library_id,source_row=source_row,peptide=sequence,aa=aa,declared_aa=declared_aa,annotation_matches_full_sequence=declared_aa==aa))
    library_book.close()
    if any(len(entries)!=19 or len({x['aa'] for x in entries})!=19 for entries in library.values()):raise ValueError('Expected 19 explicit non-parent entries per position')
    sheet=book['Figure 2C']; rows=[]
    for start,name in [(1,'A3V'),(31,'A11Va'),(61,'A11Vb'),(91,'A11Vc')]:
        if sheet.cell(start,1).value != name:
            raise ValueError('Worksheet receptor identity shifted')
        for r in range(start+10,start+28):
            aa=sheet.cell(r,2).value
            if aa not in AA or aa in 'CV':
                raise ValueError('Unexpected scan amino acid')
            for position,c in enumerate(range(3,13)):
                rows.append(response(sheet,name,r,c,position,aa,start,wt,library))
        for position,r in enumerate(range(start+10,start+20)):
            if sheet.cell(r,15).value != f'{position+7}C':
                raise ValueError('Cysteine coordinate mismatch')
            rows.append(response(sheet,name,r,16,position,'C',start,wt,library))
    book.close()
    if len(rows)!=760:raise ValueError('Unexpected raw-cell count')
    verified=[x for x in rows if x['peptide_identity_verified']]
    return dict(source_file=path.relative_to(ROOT).as_posix(),sheet='Figure 2C',source_sha256=sha(path),
                primary_peptide_library_file=library_path.relative_to(ROOT).as_posix(),primary_peptide_library_sha256=sha(library_path),
                primary_library_annotation_conflicts=[x for entries in library.values() for x in entries if not x['annotation_matches_full_sequence']],
                observations=rows,raw_cells=len(rows),verified_identity_cells=len(verified),
                quarantined_identity_cells=len(rows)-len(verified),
                literal_vs_library_order_disagreements=sum(x['literal_row_label_peptide']!=x['library_order_candidate_peptide'] for x in rows),
                literal_parent_absent_from_library_cells=sum(x['literal_interpretation_is_parent'] for x in rows),
                verified_unique_receptor_peptide_pairs=len({(x['receptor_name'],x['peptide']) for x in verified}),
                release_status='quarantined; cell-to-peptide ambiguity, author normalization conflict and construct identity bridge unresolved',
                training_release_created=False, binding_labels_created=False)

def response(sheet,name,r,c,position,aa,start,wt,library):
    raw=sheet.cell(r,c).value
    minimum=sheet.cell(start+7,c).value; maximum=sheet.cell(start+8,c).value
    if not all(isinstance(v,(int,float)) for v in [raw,minimum,maximum]) or maximum<=minimum:
        raise ValueError('Invalid raw response/controls')
    literal=wt[:position]+aa+wt[position+1:]
    entries=library[position]
    literal_entry=next((x for x in entries if x['peptide']==literal),None)
    ordered=next(x for x in entries if x['aa']=='C') if c==16 else [x for x in entries if x['aa']!='C'][r-start-10]
    verified=(c==16 or wt[position]=='V')
    if verified and (literal_entry is None or literal!=ordered['peptide']):raise ValueError('Explicit source identity does not match library')
    return dict(receptor_name=name,peptide=literal if verified else None,source_cell=sheet.cell(r,c).coordinate,
                peptide_identity_verified=verified,
                peptide_identity_status='explicit Cys ID or V-native column agrees with library' if verified else 'non-V-native column: global row labels conflict with complete non-parent library; plate map required',
                literal_row_label_peptide=literal,library_order_candidate_peptide=ordered['peptide'],
                literal_library_source_row=literal_entry['source_row'] if literal_entry else None,
                library_order_source_row=ordered['source_row'],library_order_mapping_is_assumed=not verified,
                raw_NFAT_EGFP_value=raw,minimum_Jurkat_background=minimum,maximum_PMA_ionomycin=maximum,
                minimum_cell=sheet.cell(start+7,c).coordinate,maximum_cell=sheet.cell(start+8,c).coordinate,
                derived_PMA_percent_unclipped=100*(raw-minimum)/(maximum-minimum),
                derived_endpoint='explicit curator calculation; not claimed author cognate-normalized endpoint',
                is_parent_sequence=False if verified else None,literal_interpretation_is_parent=literal==wt,
                peptide_pulse_uM=10,incubation_hours=16,
                assay='JASP90_CD8 NFAT-EGFP flow cytometry with HLA-matched K562',
                training_allowed=False,independent_acceptance_allowed=False)

def build(out):
    if out.exists():
        raise FileExistsError(out)
    sequences=sequence_ledger();cells=kras_cells()
    publication_pairs=[('PTEAM_publisher_mmc4.xlsx','pteam_rnf43.xlsx'),('PTEAM_publisher_mmc5.xlsx','pteam_cmv.xlsx')]
    pairs=[]
    for publisher,zenodo in publication_pairs:
        a=D/'sources'/publisher;b=ROOT/'docs/six_blockers_20261021/sources'/zenodo
        if sha(a)!=sha(b):raise ValueError('Publisher/Zenodo workbook differs')
        pairs.append(dict(publisher_file=a.relative_to(ROOT).as_posix(),zenodo_file=b.relative_to(ROOT).as_posix(),sha256=sha(a),byte_identical=True))
    out.mkdir()
    (out/'sequence_evidence.json').write_text(json.dumps(sequences,indent=2)+'\n',encoding='utf-8',newline='\n')
    (out/'kras_raw_cell_ledger.json').write_text(json.dumps(cells,indent=2)+'\n',encoding='utf-8',newline='\n')
    hashes={p.relative_to(ROOT).as_posix():sha(p) for p in sorted((D/'sources').iterdir()) if p.is_file()}
    (out/'source_checks.json').write_text(json.dumps(dict(publisher_Zenodo_pairs=pairs,source_sha256=hashes,training_release_created=False),indent=2)+'\n',encoding='utf-8',newline='\n')
    print(json.dumps(dict(receptors=len(sequences['receptors']),raw_cells=cells['raw_cells'],verified_identity_cells=cells['verified_identity_cells'],release_created=False)))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);build(p.parse_args().out)

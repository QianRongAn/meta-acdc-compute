"""Audit construct and normalization conflicts; no admission, fitting or test scores."""
from pathlib import Path
import argparse,hashlib,json,statistics,sys,xml.etree.ElementTree as ET
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'src')]
from scripts.curate_pteam_primary import translate

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def fasta(p,accession,species):
    lines=p.read_text().splitlines()
    if accession not in lines[0] or species not in lines[0]:raise ValueError('Constant reference identity mismatch')
    return ''.join(lines[1:])

def audit(root):
    import openpyxl
    d=root/'docs/protocol_followup_20261023';s=d/'sources'
    prior=root/'docs/six_blockers_20261021/pteam_primary_sequence_index.json'
    primary=json.loads(prior.read_text())
    human=fasta(s/'human_TRAC_uniprot.fasta','P01848','Homo sapiens')
    mouse=fasta(s/'mouse_TRAC_uniprot.fasta','P01849','Mus musculus')
    refs={'human_P01848':human,'mouse_P01849':mouse}
    xmlpath=root/'docs/six_blockers_20261021/sources/PMC11480844.xml'
    tree=ET.parse(xmlpath)
    locators={x.get('id'):' '.join(x.itertext()) for x in tree.iter('sec') if x.get('id') in ['sec4.4.2','sec4.4.3.2','sec4.4.3.3']}
    if 'hTRAC' not in locators['sec4.4.2']:raise ValueError('Methods construct claim missing')
    rows=[]
    for old in primary['receptors']:
        cmv=old['source_file']=='pteam_cmv.xlsx'
        if cmv:
            translated=translate(old['alpha']['dna'])
            aa=translated['amino_acids'];cdr=old['cdr3_alpha']
            start=aa.find('IQNP',aa.index(cdr)+len(cdr))
            if start<0:raise ValueError('Unlocated post-CDR constant anchor')
            observed=aa[start:]
            # This is a source-aligned constant segment, not mature-variable annotation.
            diff=[dict(position=i+1,reference=a,published=b) for i,(a,b) in enumerate(zip(mouse,observed)) if a!=b]
            rows.append(dict(receptor_name=old['name'],source_file=old['source_file'],source_row=old['source_row'],
                             DNA_alpha_cell='K'+str(old['source_row']),
                             sequence_scope='published contiguous chain includes leader, variable/J and constant',
                             observed_alpha_constant=observed,constant_anchor_zero_index=start,
                             methods_alpha_constant='human TRAC; ENSG00000277734.8',
                             exact_human_reference_match=observed==human,
                             mouse_reference_same_length=len(observed)==len(mouse),mouse_reference_substitutions=diff,
                             methods_sequence_conflict=True,prior_DNA_pair_passes=old['DNA_pair_passes'],
                             normalization='fraction of transgenic TCR-expressing cells; raw denominators not supplied in examined workbook',
                             peptide_pulse_uM=None,training_allowed=False,independent_acceptance_allowed=False))
        else:
            ac=translate(old['alpha_constant_DNA']);bc=translate(old['beta_constant_DNA'])
            rows.append(dict(receptor_name=old['name'],source_file=old['source_file'],source_row=old['source_row'],
                             sequence_scope='separate published variable/leader and constant DNA; no assembled cassette invented',
                             alpha_constant_translation=ac,beta_constant_translation=bc,
                             methods_sequence_conflict=None,
                             normalization='raw activation * (100 / mean PMA-ionomycin control)',
                             peptide_pulse_uM=None,prior_DNA_pair_passes=old['DNA_pair_passes'],
                             author_excluded=old['author_excluded'],training_allowed=False,independent_acceptance_allowed=False))
    cmv=[x for x in rows if x['source_file']=='pteam_cmv.xlsx']
    if len(cmv)!=20 or any(not x['mouse_reference_same_length'] or x['mouse_reference_substitutions']!=[dict(position=47,reference='T',published='C')] for x in cmv):
        raise ValueError('CMV constant conflict does not reproduce for every source row')
    sourcepath=root/'docs/measured_intake_20261021/source_metadata/kras_measurements.xlsx'
    source=openpyxl.load_workbook(sourcepath,data_only=True)
    sheet=source['Figure 2C']
    ledgerpath=d/'kras_raw_cell_ledger_v2.json'
    oldcells=json.loads(ledgerpath.read_text())
    bases=dict(A3V=1,A11Va=31,A11Vb=61,A11Vc=91);normalization=[]
    for x in oldcells['observations']:
        base=bases[x['receptor_name']];coord=x['source_cell'];col=16 if coord.startswith('P') else ord(coord[0])-ord('A')+1
        cognate=sheet.cell(base+5,col);background=sheet.cell(base+6,col)
        raw=x['raw_NFAT_EGFP_value'];lo=background.value;hi=cognate.value
        if not all(isinstance(z,(int,float)) for z in [lo,hi]) or hi<=lo:raise ValueError('Invalid cognate/background pair')
        method=100*(raw-lo)/(hi-lo);figure=x['derived_PMA_percent_unclipped']
        normalization.append(dict(receptor_name=x['receptor_name'],peptide=x['peptide'],source_cell=coord,
                                  peptide_identity_verified=x['peptide_identity_verified'],
                                  peptide_identity_status=x['peptide_identity_status'],
                                  cognate_cell=cognate.coordinate,cognate_value=hi,no_peptide_APC_cell=background.coordinate,no_peptide_APC_value=lo,
                                  methods_interpretation_unclipped=method,methods_interpretation_clipped=max(0,min(100,method)),
                                  figure_PMA_Jurkat_interpretation_unclipped=figure,
                                  figure_PMA_Jurkat_interpretation_clipped=max(0,min(100,figure)),
                                  authors_exact_implemented_normalization_verified=False,training_allowed=False))
    source.close()
    inputs=[prior,xmlpath,sourcepath,ledgerpath,root/oldcells['primary_peptide_library_file'],s/'human_TRAC_uniprot.fasta',s/'mouse_TRAC_uniprot.fasta',s/'PTEAM_full_article_and_supp.pdf']
    return dict(source_sha256={p.relative_to(root).as_posix():sha(p) for p in inputs},
                primary_methods_locators=locators,constant_references=refs,pteam_constructs=rows,
                kras_normalization_views=normalization,
                summary=dict(CMV_construct_conflicts=len(cmv),prior_sequence_pass_CMV_now_additionally_construct_blocked=sum(x['prior_DNA_pair_passes'] for x in cmv),
                             RNF_separate_segment_records=7,KRAS_raw_cells_with_both_control_views=len(normalization),
                             KRAS_identity_confirmed_cells=oldcells['verified_identity_cells'],KRAS_identity_quarantined_cells=oldcells['quarantined_identity_cells'],
                             primary_library_annotation_conflicts=len(oldcells['primary_library_annotation_conflicts']),
                             normalized_values_are_curator_interpretations=True,training_release_created=False,new_test_created=False),
                interpretation='PTEAM CMV workbook alpha constant resembles mouse P01849 with T47C, contradicting Methods human-TRAC wording. It does not establish which construct was actually used. KRAS two view calculations expose a protocol ambiguity; neither is promoted to author truth or training labels.',
                model_trained=False,model_scored=False,test_labels_used=False,partition_assignment=None)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    if a.out.exists():raise FileExistsError(a.out)
    result=audit(ROOT);a.out.write_bytes((json.dumps(result,indent=2)+'\n').encode());print(json.dumps(result['summary']))
